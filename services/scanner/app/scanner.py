"""本地 OMR 核心：定位、透视矫正、读框、判分、导出。"""
from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.constants import (AMBIGUITY_MARGIN, CONFIG_DIR, FILL_THRESHOLD, OUTPUT_DIR,
                           TEMPLATE_HEIGHT, TEMPLATE_WIDTH)
from app.roster_import import read_roster


STUDENT_RECORD_COLUMNS = ("Exam ID", "Chinese Name", "Year Level", "Branch", "Class", "Exam Session")
PART_SCORE_COLUMNS = (
    "Listening_Part1", "Listening_Part2", "Listening_Part3",
    "Written_Part1", "Written_Part2", "Written_Part3", "Written_Part4",
    "Written_Part5", "Written_Part6", "Written_Part7", "Written_Part8", "Writing",
)
SUMMARY_SEPARATOR_COLUMN = "__summary_separator__"
SECTION_LABELS = {
    "listening_part1": "Listening Part 1",
    "listening_part2": "Listening Part 2",
    "listening_part3": "Listening Part 3",
    "written_part1": "Written Part 1",
    "written_part4": "Written Part 4",
    "written_part5": "Written Part 5",
    "written_part6": "Written Part 6",
    "written_part7": "Written Part 7",
    "written_part8": "Written Part 8",
}
MANUAL_PART_LABELS = {
    "Written_Part2": "Written Part 2",
    "Written_Part3": "Written Part 3",
    "Writing": "Writing",
}
GRADER_SCORE_COLUMNS = {"Written_Part2", "Written_Part3"}
ANSWER_STATUS_LABELS = {
    "CORRECT": "正确",
    "INCORRECT": "错误",
    "BLANK": "空白",
    "AMBIGUOUS": "识别不明确",
    "SCORED": "已填分",
    "MANUAL": "已人工修正",
}


@dataclass
class MarkRead:
    value: str | None
    ratios: dict[str, float]


def load_template() -> dict[str, Any]:
    return json.loads((CONFIG_DIR / "template.json").read_text(encoding="utf-8"))


def load_answer_key() -> dict[str, dict[str, str]]:
    return json.loads((CONFIG_DIR / "answer_key.json").read_text(encoding="utf-8"))


def load_students() -> dict[str, dict[str, str]]:
    imported = read_roster(CONFIG_DIR / "student_list.xlsx")
    source_path = CONFIG_DIR / "roster_source.xlsx"
    original = read_roster(source_path) if source_path.is_file() else imported
    exam_position = dict(original.source_positions)["Exam ID"]
    original_ids = {
        row["Exam ID"]: str(source[exam_position]).strip()
        for row, source in zip(original.rows, original.source_rows)
    }
    return {
        row["Exam ID"]: {**row, "Original Exam ID": original_ids.get(row["Exam ID"], row["Exam ID"])}
        for row in imported.rows
    }


def student_record(student: dict[str, str] | None, exam_id: str) -> dict[str, str]:
    record = {
        "Exam ID": exam_id,
        "Chinese Name": student.get("Chinese Name", "") if student else "",
        "Year Level": student.get("Year Level", "") if student else "",
        "Branch": student.get("Branch", "") if student else "",
        "Class": student.get("Class", "") if student else "",
        "Exam Session": student.get("Exam Session", "") if student else "",
    }
    if student and student.get("Original Exam ID"):
        record["Original Exam ID"] = student["Original Exam ID"]
    return record


def update_score_totals(record: dict[str, Any]) -> None:
    """从 12 个 Part 得分统一重算全卷总分。"""
    scores = {column: int(record.get(column) or 0) for column in PART_SCORE_COLUMNS}
    record["Total"] = sum(scores.values())
    record.pop("Listening Total", None)
    record.pop("Written Total", None)


def update_review_flag(record: dict[str, Any], item_rows: list[dict[str, Any]]) -> None:
    """综合判定是否需复核。

    单题空白只保留在逐题明细；整 Part 空白、考号无法识别或考号不在
    名单中时，必须进入复核队列。
    """
    source = record.get("Source Image", "")
    relevant = [row for row in item_rows if row.get("Source Image") == source]
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in relevant:
        grouped.setdefault(str(row.get("Section", "")), []).append(row)

    blank_parts: list[str] = []
    blank_details: list[str] = []
    unresolved_blank_parts: list[str] = []
    reviewed_sections = set(record.get("Reviewed Blank Sections", []))
    for section, rows in grouped.items():
        statuses = [
            row.get("Answer Status") or ("BLANK" if not row.get("Marked Answer") else "ANSWERED")
            for row in rows
        ]
        if rows and all(status == "BLANK" for status in statuses):
            label = SECTION_LABELS.get(section, section)
            numbers = "、".join(str(row.get("Question", "")) for row in rows)
            blank_parts.append(label)
            blank_details.append(f"{label}：{numbers}")
            if section not in reviewed_sections:
                unresolved_blank_parts.append(label)

    entry_states = record.get("Score Entry States", {})
    reviewed_score_fields = set(record.get("Reviewed Blank Scores", []))
    unread_grader_scores: list[str] = []
    for column, label in MANUAL_PART_LABELS.items():
        state = entry_states.get(column)
        if column in GRADER_SCORE_COLUMNS and state in {"BLANK", "AMBIGUOUS"}:
            unread_grader_scores.append(label)
        elif state == "BLANK":
            blank_parts.append(label)
            blank_details.append(f"{label}：计分区域空白")
            if column not in reviewed_score_fields:
                unresolved_blank_parts.append(label)

    record["Blank Parts"] = "、".join(blank_parts)
    record["Blank Questions"] = "；".join(blank_details)
    reasons: list[str] = []
    identity_issue = str(record.get("Identity Issue", "")).strip()
    if identity_issue:
        reasons.append(identity_issue)
    if unread_grader_scores:
        reasons.append(f"教师登分区未成功读取：{'、'.join(unread_grader_scores)}")
    if unresolved_blank_parts:
        reasons.append("答题卡整Part未填写")
    record["Needs Review"] = "是" if reasons else "否"
    record["Review Reason"] = "；".join(reasons)


def _read_image(path: Path) -> np.ndarray:
    if path.suffix.lower() in {".heic", ".heif"}:
        try:
            from pillow_heif import register_heif_opener
            from PIL import Image
            register_heif_opener()
            return cv2.cvtColor(np.array(Image.open(path).convert("RGB")), cv2.COLOR_RGB2BGR)
        except ImportError as exc:
            raise ValueError("HEIC support requires pillow-heif") from exc
    image = cv2.imread(str(path))
    if image is None:
        raise ValueError("Unreadable image")
    return image


def _find_registration_marks(image: np.ndarray) -> np.ndarray:
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    binary = cv2.threshold(gray, 55, 255, cv2.THRESH_BINARY_INV)[1]
    contours, _ = cv2.findContours(binary, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    height, width = gray.shape
    candidates: list[tuple[float, float]] = []
    min_area, max_area = width * height * 0.00015, width * height * 0.02
    for contour in contours:
        x, y, w, h = cv2.boundingRect(contour)
        area = w * h
        if min_area <= area <= max_area and 0.65 <= w / h <= 1.35:
            candidates.append((x + w / 2, y + h / 2))
    if len(candidates) < 4:
        raise ValueError("Could not find four registration marks")
    targets = [(0, 0), (width, 0), (width, height), (0, height)]
    selected = []
    used: set[int] = set()
    for tx, ty in targets:
        options = sorted(((cx - tx) ** 2 + (cy - ty) ** 2, i) for i, (cx, cy) in enumerate(candidates) if i not in used)
        if not options:
            raise ValueError("Could not uniquely assign registration marks")
        _, index = options[0]
        used.add(index)
        selected.append(candidates[index])
    return np.float32(selected)


def normalize(image: np.ndarray, template: dict[str, Any]) -> np.ndarray:
    """将手机照片通过四个黑色定位块映射到标准答题卡像素坐标。"""
    source = _find_registration_marks(image)
    marks = template["registration_marks"]
    destination = np.float32([[marks[name]["cx"], marks[name]["cy"]] for name in ("top_left", "top_right", "bottom_right", "bottom_left")])
    transform = cv2.getPerspectiveTransform(source, destination)
    return cv2.warpPerspective(image, transform, (TEMPLATE_WIDTH, TEMPLATE_HEIGHT), borderValue=(255, 255, 255))


def _fill_ratio(image: np.ndarray, box: dict[str, Any]) -> float:
    x, y, w, h = (int(box[key]) for key in ("x", "y", "w", "h"))
    inset = max(3, int(min(w, h) * 0.22))
    crop = cv2.cvtColor(image[y + inset:y + h - inset, x + inset:x + w - inset], cv2.COLOR_BGR2GRAY)
    if crop.size == 0:
        return 0.0
    return float(np.mean(crop < 140))


def _read_one(image: np.ndarray, options: dict[str, dict[str, Any]]) -> MarkRead:
    ratios = {label: _fill_ratio(image, box) for label, box in options.items()}
    ordered = sorted(ratios.items(), key=lambda item: item[1], reverse=True)
    best_label, best = ordered[0]
    second = ordered[1][1] if len(ordered) > 1 else 0.0
    if best < FILL_THRESHOLD or best - second < AMBIGUITY_MARGIN:
        return MarkRead(None, ratios)
    return MarkRead(best_label, ratios)


def _read_questions(image: np.ndarray, groups: dict[str, dict[str, dict[str, Any]]]) -> tuple[dict[str, str | None], dict[str, str], bool]:
    answers, states, ambiguous = {}, {}, False
    for number, options in groups.items():
        read = _read_one(image, options)
        answers[number] = read.value
        states[number] = "ANSWERED" if read.value is not None else "BLANK" if max(read.ratios.values()) < FILL_THRESHOLD else "AMBIGUOUS"
        ambiguous = ambiguous or read.value is None
    return answers, states, ambiguous


def _score_answers(answers: dict[str, str | None], key: dict[str, str], states: dict[str, str]) -> tuple[int, list[dict[str, Any]]]:
    items, score = [], 0
    for number, correct in key.items():
        marked = answers.get(number)
        is_correct = marked == correct
        score += int(is_correct)
        answer_status = states.get(number, "BLANK")
        if answer_status == "ANSWERED":
            answer_status = "CORRECT" if is_correct else "INCORRECT"
        items.append({
            "Question": number,
            "Original Marked Answer": marked or "",
            "Marked Answer": marked or "",
            "Correct Answer": correct,
            "Original Item Score": 1 if is_correct else 0,
            "Is Correct": "Y" if is_correct else "N",
            "Original Answer Status": answer_status,
            "Answer Status": answer_status,
            "Manual Correction": "否",
            "Review Basis": "",
            "Operator": "",
            "Changed At": "",
        })
    return score, items


def scan_image(path: Path, template: dict[str, Any] | None = None, students: dict[str, dict[str, str]] | None = None) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    template = template or load_template()
    students = students or load_students()
    key = load_answer_key()
    normalized = normalize(_read_image(path), template)
    digits, id_bad = [], False
    for position in map(str, range(1, 7)):
        read = _read_one(normalized, template["exam_id"][position])
        digits.append(read.value or "?")
        id_bad = id_bad or read.value is None
    exam_id = "".join(digits)
    answers_by_section: dict[str, dict[str, str | None]] = {}
    states_by_section: dict[str, dict[str, str]] = {}
    mark_bad = False
    for source_key, section in [("part1", "listening_part1"), ("part2", "listening_part2"), ("part3", "listening_part3")]:
        answers, states, bad = _read_questions(normalized, template["listening"][source_key])
        answers_by_section[section] = answers
        states_by_section[section] = states
        mark_bad = mark_bad or bad
    for source_key, section in [("part1", "written_part1"), ("part4", "written_part4"), ("part5", "written_part5"), ("part6", "written_part6"), ("part7", "written_part7"), ("part8", "written_part8")]:
        answers, states, bad = _read_questions(normalized, template["written_objective"][source_key])
        answers_by_section[section] = answers
        states_by_section[section] = states
        mark_bad = mark_bad or bad
    score_values, score_entry_states, score_bad = {}, {}, False
    for template_key, label in [("score_part2", "Written_Part2"), ("score_part3", "Written_Part3"), ("score_writing", "Writing")]:
        read = _read_one(normalized, template[template_key])
        score_values[label] = int(read.value) if read.value is not None else None
        score_entry_states[label] = "SCORED" if read.value is not None else "BLANK" if max(read.ratios.values()) < FILL_THRESHOLD else "AMBIGUOUS"
        score_bad = score_bad or read.value is None
    mark_bad = mark_bad or score_bad
    section_scores: dict[str, int] = {}
    item_rows: list[dict[str, Any]] = []
    for section, answers in answers_by_section.items():
        section_scores[section], rows = _score_answers(answers, key[section], states_by_section[section])
        for row in rows:
            row.update({"Exam ID": exam_id, "Section": section, "Source Image": path.name})
        item_rows.extend(rows)
    roster = students.get(exam_id)
    identity_issue = "考号识别异常" if id_bad else "考号不在名单中" if roster is None else ""
    status = "CHECK_ID" if identity_issue else "CHECK_MARK" if mark_bad else "OK"
    record = student_record(roster, exam_id)
    record.update({
        "Listening_Part1": section_scores["listening_part1"],
        "Listening_Part2": section_scores["listening_part2"],
        "Listening_Part3": section_scores["listening_part3"],
        "Written_Part1": section_scores["written_part1"],
        "Written_Part4": section_scores["written_part4"],
        "Written_Part5": section_scores["written_part5"],
        "Written_Part6": section_scores["written_part6"],
        "Written_Part7": section_scores["written_part7"],
        "Written_Part8": section_scores["written_part8"],
    })
    record.update(score_values)
    record["Original Scores"] = {
        **{column: int(record.get(column) or 0) for column in PART_SCORE_COLUMNS if column not in score_values},
        **score_values,
    }
    record["Score Entry States"] = score_entry_states
    record["Original Score Entry States"] = dict(score_entry_states)
    update_score_totals(record)
    record["Status"] = status
    record["File Name"] = path.name
    record["Source Image"] = path.name
    record["Identity Issue"] = identity_issue
    update_review_flag(record, item_rows)
    return record, item_rows


RESULT_COLUMNS = [
    "File Name", *STUDENT_RECORD_COLUMNS, *PART_SCORE_COLUMNS, "Total",
    SUMMARY_SEPARATOR_COLUMN, "Blank Parts", "Blank Questions", "Needs Review",
    "Review Reason", "Identity Issue",
]
RESULT_HEADER_LABELS = {
    "Exam ID": "考号",
    "Chinese Name": "中文名",
    "Year Level": "年级",
    "Branch": "分校",
    "Class": "班级",
    "Exam Session": "笔试时间",
    SUMMARY_SEPARATOR_COLUMN: "",
    "Blank Parts": "空白Part",
    "Blank Questions": "空白题号",
    "Needs Review": "需要人工核对学生试卷",
    "Review Reason": "异常原因",
    "Identity Issue": "身份异常",
}

ITEM_COLUMNS = [
    *STUDENT_RECORD_COLUMNS, "Part", "Question", "Original Marked Answer",
    "Marked Answer", "Correct Answer", "Original Item Score", "Item Score",
    "Original Answer Status", "Answer Status", "Manual Correction", "Review Basis",
    "Operator", "Changed At", "Part Status", "Source Image",
]
ITEM_HEADER_LABELS = {
    **RESULT_HEADER_LABELS,
    "Part": "Part",
    "Question": "题号",
    "Original Marked Answer": "原始识别答案",
    "Marked Answer": "最终生效答案",
    "Correct Answer": "正确答案",
    "Original Item Score": "原始得分",
    "Item Score": "得分",
    "Original Answer Status": "原始作答状态",
    "Answer Status": "作答状态",
    "Manual Correction": "是否人工修正",
    "Review Basis": "核查依据",
    "Operator": "操作人",
    "Changed At": "修改时间",
    "Part Status": "Part状态",
    "Source Image": "答题卡文件",
}
REVIEW_COLUMNS = [
    *STUDENT_RECORD_COLUMNS, *PART_SCORE_COLUMNS, "Total", "Source Image",
    "Blank Parts", "Blank Questions", "Needs Review", "Review Reason", "Review Status",
    "Instructions",
]
REVIEW_HEADER_LABELS = {
    **RESULT_HEADER_LABELS,
    "Source Image": "答题卡文件",
    "Blank Parts": "空白Part",
    "Blank Questions": "空白题号",
    "Review Reason": "异常原因",
    "Review Status": "处理状态",
    "Instructions": "说明",
}


def _style_sheet(ws, header_color: str = "1F4E78") -> None:
    header = PatternFill("solid", fgColor=header_color)
    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment = header, Font(color="FFFFFF", bold=True), Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for column in ws.columns:
        width = min(28, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        ws.column_dimensions[column[0].column_letter].width = width
    exam_id_column = next((cell.column for cell in ws[1] if cell.value in {"Exam ID", "考号"}), None)
    if exam_id_column is not None:
        for row in ws.iter_rows(min_row=2, min_col=exam_id_column, max_col=exam_id_column):
            row[0].number_format = "@"


def _export_item_rows(records: list[dict[str, Any]], item_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    records_by_source = {record.get("Source Image", ""): record for record in records}
    part_rows: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for row in item_rows:
        part_rows.setdefault((str(row.get("Source Image", "")), str(row.get("Section", ""))), []).append(row)

    exported: list[dict[str, Any]] = []
    for row in item_rows:
        source = str(row.get("Source Image", ""))
        record = records_by_source.get(source, {})
        section = str(row.get("Section", ""))
        peers = part_rows.get((source, section), [])
        whole_part_blank = bool(peers) and all((peer.get("Answer Status") or ("BLANK" if not peer.get("Marked Answer") else "ANSWERED")) == "BLANK" for peer in peers)
        exported.append({
            **{column: record.get(column, "") for column in STUDENT_RECORD_COLUMNS},
            "Part": SECTION_LABELS.get(section, section),
            "Question": row.get("Question", ""),
            "Original Marked Answer": row.get("Original Marked Answer", row.get("Marked Answer", "")),
            "Marked Answer": row.get("Marked Answer", ""),
            "Correct Answer": row.get("Correct Answer", ""),
            "Original Item Score": row.get("Original Item Score", 1 if row.get("Is Correct") == "Y" else 0),
            "Item Score": 1 if row.get("Is Correct") == "Y" else 0,
            "Original Answer Status": ANSWER_STATUS_LABELS.get(
                row.get("Original Answer Status", row.get("Answer Status", "")),
                row.get("Original Answer Status", row.get("Answer Status", "")),
            ),
            "Answer Status": ANSWER_STATUS_LABELS.get(
                row.get("Answer Status") or ("BLANK" if not row.get("Marked Answer") else "CORRECT" if row.get("Is Correct") == "Y" else "INCORRECT"),
                row.get("Answer Status", ""),
            ),
            "Manual Correction": row.get("Manual Correction", "否"),
            "Review Basis": row.get("Review Basis", ""),
            "Operator": row.get("Operator", ""),
            "Changed At": row.get("Changed At", ""),
            "Part Status": "整Part空白" if whole_part_blank else "正常",
            "Source Image": source,
        })

    for record in records:
        states = record.get("Score Entry States", {})
        original_states = record.get("Original Score Entry States", states)
        original_scores = record.get("Original Scores", {})
        score_audits = record.get("Manual Score Audits", {})
        for column, label in MANUAL_PART_LABELS.items():
            state = states.get(column, "MANUAL" if record.get("Manual Override") else "SCORED")
            original_state = original_states.get(column, state)
            raw_original_score = original_scores.get(column, record.get(column))
            original_score = "" if raw_original_score is None else int(raw_original_score)
            final_score = int(record.get(column) or 0)
            manually_corrected = state == "MANUAL" or original_score != final_score
            score_audit = score_audits.get(column, {}) if state == "MANUAL" else {}
            exported.append({
                **{field: record.get(field, "") for field in STUDENT_RECORD_COLUMNS},
                "Part": label,
                "Question": "",
                "Original Marked Answer": original_score,
                "Marked Answer": final_score,
                "Correct Answer": "",
                "Original Item Score": original_score,
                "Item Score": final_score,
                "Original Answer Status": (
                    "教师登分区未成功读取"
                    if column in GRADER_SCORE_COLUMNS and original_state in {"BLANK", "AMBIGUOUS"}
                    else ANSWER_STATUS_LABELS.get(original_state, original_state)
                ),
                "Answer Status": ANSWER_STATUS_LABELS.get(state, state),
                "Manual Correction": "是" if manually_corrected else "否",
                "Review Basis": score_audit.get("Review Basis", record.get("Last Review Basis", "") if state == "MANUAL" else ""),
                "Operator": score_audit.get("Operator", record.get("Last Operator", "") if state == "MANUAL" else ""),
                "Changed At": score_audit.get("Changed At", record.get("Last Changed At", "") if state == "MANUAL" else ""),
                "Part Status": (
                    "教师登分区未成功读取"
                    if column in GRADER_SCORE_COLUMNS and state in {"BLANK", "AMBIGUOUS"}
                    else "整Part空白" if state == "BLANK" else "正常"
                ),
                "Source Image": record.get("Source Image", ""),
            })
    return exported


def _review_rows(records: list[dict[str, Any]], audit_log: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Sheet3 只列出系统内人工修正后仍无法解决的问题。

    人工修改前后值、操作人和时间已由 Sheet2 保留，不在 Sheet3 重复。
    audit_log 参数仅为保持现有导出调用兼容。
    """
    del audit_log
    rows: list[dict[str, Any]] = []
    for record in records:
        source = str(record.get("Source Image", ""))
        has_blank_part = bool(str(record.get("Blank Parts", "")).strip())
        has_unresolved_issue = record.get("Needs Review") == "是"
        if not has_blank_part and not has_unresolved_issue:
            continue
        reason = str(record.get("Review Reason", "")).strip()
        if has_blank_part:
            reason = "整Part未填涂，需复查原试卷"
        rows.append({
            **{column: record.get(column, "") for column in STUDENT_RECORD_COLUMNS},
            **{column: record.get(column, "") for column in PART_SCORE_COLUMNS},
            "Total": record.get("Total", ""),
            "Source Image": source,
            "Blank Parts": record.get("Blank Parts", ""),
            "Blank Questions": record.get("Blank Questions", ""),
            "Needs Review": "是",
            "Review Reason": reason,
            "Review Status": "待复查原卷" if has_blank_part else "待处理",
            "Instructions": "本次学生整Part未填涂，建议复查原试卷。" if has_blank_part else "该问题无法在当前图片复核流程中解决，请单独核查。",
        })
    return rows


def _append_sheet(workbook: Workbook, title: str, columns: list[str], rows: list[dict[str, Any]], labels: dict[str, str], header_color: str = "1F4E78"):
    worksheet = workbook.create_sheet(title)
    worksheet.append([labels.get(column, column) for column in columns])
    for row in rows:
        worksheet.append([row.get(column, "") for column in columns])
    _style_sheet(worksheet, header_color)
    return worksheet


def _append_summary_sheet(workbook: Workbook, records: list[dict[str, Any]]):
    """使用本次导入的原始名单列，在其后追加扫描成绩。

    旧用户数据没有 roster_source.xlsx 时自动回退到原六列汇总格式。
    """
    source_path = CONFIG_DIR / "roster_source.xlsx"
    if not source_path.is_file():
        # Preserve a known original student number even in the legacy format.
        exported_records = [
            {**record, "Exam ID": record.get("Original Exam ID") or record.get("Exam ID", "")}
            for record in records
        ]
        return _append_sheet(workbook, "成绩汇总", RESULT_COLUMNS, exported_records, RESULT_HEADER_LABELS)

    imported = read_roster(source_path)
    source_headers = list(imported.source_headers)
    positions = dict(imported.source_positions)
    source_by_id = {
        row["Exam ID"].zfill(6): list(source_row)
        for row, source_row in zip(imported.rows, imported.source_rows)
    }
    append_columns = [
        SUMMARY_SEPARATOR_COLUMN, "File Name", *PART_SCORE_COLUMNS, "Total",
        "Blank Parts", "Blank Questions", "Needs Review", "Review Reason", "Identity Issue",
    ]
    worksheet = workbook.create_sheet("成绩汇总")
    worksheet.append(source_headers + [RESULT_HEADER_LABELS.get(column, column) for column in append_columns])
    for record in records:
        exam_id = str(record.get("Exam ID", ""))
        source_row = list(source_by_id.get(exam_id.zfill(6), []))
        source_row_missing = not source_row
        if source_row_missing:
            source_row = [""] * len(source_headers)
            for canonical in STUDENT_RECORD_COLUMNS:
                position = positions.get(canonical)
                if position is not None and position < len(source_row):
                    source_row[position] = record.get(canonical, "")
        elif len(source_row) != len(source_headers):
            # 原始 Excel 可能省略行末空单元格，也可能在表头范围之外残留
            # 格式或旧数据。追加成绩前必须严格对齐表头宽度，否则复核列
            # 会整体错位，结果管理会把 Part 分数误读成身份异常。
            source_row = source_row[:len(source_headers)]
            source_row.extend([""] * (len(source_headers) - len(source_row)))
        exam_position = positions.get("Exam ID")
        # 名单内学生保留原始 S + 5 位学号；阅卷匹配只在内部
        # 使用 0 + 5 位考号。只有名单外记录才需要回填内部考号。
        if source_row_missing and exam_position is not None and exam_position < len(source_row):
            source_row[exam_position] = record.get("Original Exam ID") or exam_id
        worksheet.append(source_row + [record.get(column, "") for column in append_columns])
    _style_sheet(worksheet)
    return worksheet


def export_results(records: list[dict[str, Any]], item_rows: list[dict[str, Any]], output_dir: Path = OUTPUT_DIR, export_label: str | None = None, include_items: bool = True, audit_log: list[dict[str, Any]] | None = None) -> tuple[Path, Path | None]:
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d%H%M")
    result_stem = export_label or "results"
    suffix = ""
    attempt = 1
    while True:
        result_path = output_dir / f"{result_stem}_{timestamp}{suffix}.xlsx"
        if not result_path.exists():
            break
        attempt += 1
        suffix = f"_{attempt:02d}"
    for record in records:
        update_score_totals(record)
        update_review_flag(record, item_rows)
    ordered_records = sorted(
        records,
        key=lambda record: (bool(record.get("Identity Issue") or not record.get("Chinese Name")),),
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    _append_summary_sheet(workbook, ordered_records)
    if include_items:
        detail_sheet = _append_sheet(workbook, "逐题明细", ITEM_COLUMNS, _export_item_rows(records, item_rows), ITEM_HEADER_LABELS, "4472C4")
        review_sheet = _append_sheet(workbook, "需复核名单", REVIEW_COLUMNS, _review_rows(records, audit_log or []), REVIEW_HEADER_LABELS, "C65911")
        instructions_column = REVIEW_COLUMNS.index("Instructions") + 1
        review_sheet.column_dimensions[get_column_letter(instructions_column)].width = 48
        for row_number in range(2, review_sheet.max_row + 1):
            review_sheet.row_dimensions[row_number].height = 48
        for worksheet in (detail_sheet, review_sheet):
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
    workbook.save(result_path)
    return result_path, None


def process_folder(folder: Path, output_dir: Path = OUTPUT_DIR) -> tuple[list[dict[str, Any]], list[dict[str, Any]], tuple[Path, Path | None]]:
    template, students = load_template(), load_students()
    records, item_rows = [], []
    for path in sorted(folder.iterdir()):
        if path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".heic", ".heif"}:
            continue
        try:
            record, items = scan_image(path, template, students)
        except Exception:
            record = {column: "" for column in RESULT_COLUMNS}
            record.update({"File Name": path.name, "Status": "CHECK_MARK", "Source Image": path.name})
            items = []
        records.append(record)
        item_rows.extend(items)
    return records, item_rows, export_results(records, item_rows, output_dir)
