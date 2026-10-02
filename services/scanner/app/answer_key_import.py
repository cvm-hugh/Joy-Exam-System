"""从模板包自动生成、导入并校验结构化正确答案表。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.models import AnswerItem, TemplatePackage, ValidationError, validate_answer_key


ANSWER_COLUMNS = ("Item ID", "Question", "Score Section", "Allowed Choices", "Correct Answer")


def generate_answer_key_template(template: TemplatePackage, destination: Path) -> Path:
    """生成可供教师填写的答案 Excel；题号/选项/分项来自模板包。"""
    template.validate()
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Answer Key"
    worksheet.append(ANSWER_COLUMNS)
    for item in template.answer_items:
        worksheet.append((item.item_id, item.label, item.score_section, " / ".join(item.choices), ""))
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column, width in {"A": 18, "B": 20, "C": 22, "D": 24, "E": 18}.items():
        worksheet.column_dimensions[column].width = width
    workbook.save(destination)
    return destination


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def read_answer_key(path: Path, template: TemplatePackage) -> dict[str, str]:
    """读取答案表，并确保它与所选模板包完全一致。"""
    if path.suffix.lower() != ".xlsx":
        raise ValidationError("正确答案必须使用系统生成或同结构的 .xlsx 答案表。")
    if not path.is_file():
        raise ValidationError(f"找不到正确答案文件：{path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError(f"无法读取正确答案：{exc}") from exc
    worksheet = workbook["Answer Key"] if "Answer Key" in workbook.sheetnames else workbook.active
    header = tuple(_text(value) for value in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
    if header[: len(ANSWER_COLUMNS)] != ANSWER_COLUMNS:
        raise ValidationError("正确答案表格式不正确；请使用系统生成的答案录入模板。")
    answers: dict[str, str] = {}
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        item_id = _text(values[0] if len(values) > 0 else "")
        answer = _text(values[4] if len(values) > 4 else "")
        if item_id:
            answers[item_id] = answer
    workbook.close()
    validate_answer_key(template.answer_items, answers)
    return answers


def answer_items_from_rows(rows: list[dict[str, object]]) -> tuple[AnswerItem, ...]:
    """模板制作器写入题目定义时可复用的转换函数。"""
    return tuple(
        AnswerItem(
            item_id=str(row["item_id"]),
            label=str(row["label"]),
            choices=tuple(str(choice) for choice in row["choices"]),
            score_section=str(row["score_section"]),
        )
        for row in rows
    )

