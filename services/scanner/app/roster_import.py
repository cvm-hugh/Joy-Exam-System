"""V3 学生名单导入：保留每次考试的名单快照，不覆盖旧名单。"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
from shutil import copy2

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font, PatternFill

from app.models import ValidationError


ROSTER_COLUMNS = ("Exam ID", "Chinese Name", "Year Level", "Branch", "Class", "Exam Session")
REQUIRED_COLUMNS = ("Exam ID", "Chinese Name")
ROSTER_TEMPLATE_HEADERS = (
    "学号", "中英文名", "姓名", "年级", "上课科目", "课程", "班级", "毕业时间", "上课校区",
    "任课老师", "笔试时间", "备注", "是否报名TM", "联系电话", "精修笔试通过", "",
    "听力理解得分系数", "词汇运用得分系数", "语法运用得分系数", "交际能力得分系数",
    "阅读理解得分系数", "写作能力得分系数", "", "", "", "",
)
BRANCH_OPTIONS = (
    "牡丹广场分校", "滨河分校", "凯旋路分校", "太康路分校", "英才路分校",
    "科大分校", "纱厂路分校", "长兴街分校", "新街分校",
)
GRADE_OPTIONS = tuple(f"{name}年级" for name in "一二三四五六七八九")

# 日常名单常来自教务或参考成绩表；接受常用中文表头，但写入 App 的副本
# 始终使用标准英文列，供现有扫描核心稳定读取。
HEADER_ALIASES = {
    "Exam ID": ("Exam ID", "考号", "学号", "考生号", "准考证号", "考试号"),
    "Chinese Name": ("Chinese Name", "Name", "中文名", "姓名", "中英文名"),
    "Year Level": ("Year Level", "Grade", "年级"),
    "Branch": ("Branch", "分校", "校区", "上课校区"),
    "Class": ("Class", "班级", "班型"),
    "Exam Session": ("Exam Session", "Exam Batch", "考试时间", "笔试时间", "考试批次", "考试场次", "参考时间", "参考场次", "批次"),
}


@dataclass(frozen=True)
class RosterImport:
    source: Path
    rows: tuple[dict[str, str], ...]
    source_headers: tuple[str, ...] = ()
    source_rows: tuple[tuple[object, ...], ...] = ()
    source_positions: tuple[tuple[str, int], ...] = ()
    source_sheet: str = ""
    skipped_rows: tuple[str, ...] = ()

    @property
    def count(self) -> int:
        return len(self.rows)


def _as_text(value: object) -> str:
    if value is None:
        return ""
    # Excel may have converted an unquoted numeric ID to a float.
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header_key(value: object) -> str:
    """忽略表头中的空格、下划线与大小写，便于匹配常见导出表。"""
    return "".join(_as_text(value).replace("_", "").split()).casefold()


def _column_positions(header: list[str]) -> dict[str, int]:
    positions = {_header_key(value): index for index, value in enumerate(header) if value}
    found: dict[str, int] = {}
    for canonical, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            position = positions.get(_header_key(alias))
            if position is not None:
                found[canonical] = position
                break
    return found


def _display_name(value: str, source_header: str) -> str:
    """原始模板的“姓名”可能同时包含英文名；界面和结果查询只使用中文名。

    原 Excel 中的姓名单元格不改写，这里只生成 App 内部显示值。
    """
    if _header_key(source_header) in {
        _header_key("姓名"), _header_key("中英文名"), _header_key("Name")
    }:
        chinese = "".join(re.findall(r"[\u3400-\u9fff·•]", value))
        return chinese or value
    return value


def normalize_exam_id(value: object) -> str:
    """教务原始名单用 S + 5 位数字，答题卡填涂的是 0 + 5 位数字。

    原 Excel 单元格仍原样保留；只对 App 内部匹配键进行转换。
    """
    text = _as_text(value)
    match = re.fullmatch(r"[Ss]([0-9]{5})", text)
    if match:
        return "0" + match.group(1)
    if text.isdigit() and len(text) <= 6:
        return text.zfill(6)
    return text


def _find_roster_sheet(workbook):
    """优先选择包含名单字段最多的工作表，而不是机械读取第一个工作表。"""
    candidates = []
    for worksheet in workbook.worksheets:
        # 部分 Excel/WPS 文件的 worksheet dimension 元数据错误地只写成 A1:A1。
        # Excel 界面仍会显示全部单元格，但 openpyxl 只读模式会因此截断到第一列。
        # 忽略该元数据并按实际存储的单元格重新扫描。
        reset_dimensions = getattr(worksheet, "reset_dimensions", None)
        if callable(reset_dimensions):
            reset_dimensions()
        header = [_as_text(cell) for cell in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())]
        positions = _column_positions(header)
        candidates.append((len(positions), worksheet, header, positions))
    _, worksheet, header, positions = max(candidates, key=lambda candidate: candidate[0])
    return worksheet, header, positions


def read_roster(path: Path) -> RosterImport:
    if path.suffix.lower() != ".xlsx":
        raise ValidationError("学生名单必须是 .xlsx 文件。")
    if not path.is_file():
        raise ValidationError(f"找不到学生名单：{path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl exception hierarchy is version-dependent
        raise ValidationError(f"无法读取学生名单：{exc}") from exc
    worksheet, header, positions = _find_roster_sheet(workbook)
    source_sheet = worksheet.title
    missing = [column for column in REQUIRED_COLUMNS if column not in positions]
    if missing:
        visible_headers = "、".join(header[:12]) or "（空表）"
        workbook.close()
        raise ValidationError(
            f"学生名单缺少必需列：{'、'.join(missing)}。"
            f"已检查工作表“{worksheet.title}”，识别到的表头：{visible_headers}。"
        )
    rows: list[dict[str, str]] = []
    source_rows: list[tuple[object, ...]] = []
    exam_ids: set[str] = set()
    errors: list[str] = []
    skipped_rows: list[str] = []
    name_position = positions.get("Chinese Name")
    name_aliases = {_header_key(alias) for alias in HEADER_ALIASES["Chinese Name"]}
    name_fallback_positions = [
        index for index, source_header in enumerate(header)
        if index != name_position and _header_key(source_header) in name_aliases
    ]
    for row_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), 2):
        row = {
            column: _as_text(values[positions[column]] if column in positions and positions[column] < len(values) else "")
            for column in ROSTER_COLUMNS
        }
        exam_position = positions.get("Exam ID")
        if exam_position is not None:
            row["Exam ID"] = normalize_exam_id(row["Exam ID"])
        if name_position is not None:
            row["Chinese Name"] = _display_name(
                row["Chinese Name"],
                header[name_position] if name_position < len(header) else "",
            )
        # 第三列“姓名”有时是 Excel 公式；文件没有保存公式缓存时，
        # openpyxl(data_only=True) 会读到空值。此时从同行“中英文名”
        # 提取中文姓名，不要误报整行缺少中文名。
        if not row["Chinese Name"]:
            for fallback_position in name_fallback_positions:
                if fallback_position >= len(values):
                    continue
                fallback_name = _as_text(values[fallback_position])
                if fallback_name:
                    row["Chinese Name"] = _display_name(
                        fallback_name, header[fallback_position]
                    )
                    break
        if not any(row.values()):
            continue
        if not row["Exam ID"]:
            skipped_rows.append(
                f"第 {row_number} 行（{row['Chinese Name'] or '未命名'}）缺少学号"
            )
            continue
        if not row["Chinese Name"]:
            errors.append(f"第 {row_number} 行缺少中文名")
            continue
        if row["Exam ID"] in exam_ids:
            errors.append(f"第 {row_number} 行 Exam ID 重复：{row['Exam ID']}")
            continue
        exam_ids.add(row["Exam ID"])
        rows.append(row)
        source_rows.append(tuple(values))
    workbook.close()
    if errors:
        raise ValidationError("学生名单校验失败：" + "；".join(errors))
    if not rows:
        raise ValidationError("学生名单没有有效学生记录。")
    return RosterImport(
        source=path,
        rows=tuple(rows),
        source_headers=tuple(header),
        source_rows=tuple(source_rows),
        source_positions=tuple(positions.items()),
        source_sheet=source_sheet,
        skipped_rows=tuple(skipped_rows),
    )


def save_roster_snapshot(
    source: Path,
    destination_directory: Path,
    snapshot_name: str = "student_list.xlsx",
    *,
    overwrite: bool = False,
) -> Path:
    """保存已验证名单的标准化副本，供现有扫描核心稳定读取。"""
    imported = read_roster(source)
    destination_directory.mkdir(parents=True, exist_ok=True)
    destination = destination_directory / snapshot_name
    if destination.exists() and not overwrite:
        raise FileExistsError(f"名单快照已存在：{destination}")
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Students"
    worksheet.append(ROSTER_COLUMNS)
    for row in imported.rows:
        worksheet.append([row[column] for column in ROSTER_COLUMNS])
    for column in worksheet.columns:
        worksheet.column_dimensions[column[0].column_letter].width = max(
            12, min(30, max(len(str(cell.value or "")) for cell in column) + 2)
        )
    worksheet.freeze_panes = "A2"
    workbook.save(destination)
    source_snapshot = destination_directory / "roster_source.xlsx"
    if source.resolve() != source_snapshot.resolve():
        copy2(source, source_snapshot)
    return destination


def generate_roster_template(destination: Path) -> Path:
    """生成与教务“原始名单模板”同列的空白表。"""
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.append(ROSTER_TEMPLATE_HEADERS)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = "A1:Z1"
    branch_validation = DataValidation(
        type="list",
        formula1='"' + ",".join(BRANCH_OPTIONS) + '"',
        allow_blank=False,
        error="请选择下拉列表中的完整分校名称。",
        errorTitle="分校名称不在名单中",
        prompt="请选择学生所属分校。",
        promptTitle="选择分校",
        showErrorMessage=True,
        showInputMessage=True,
    )
    grade_validation = DataValidation(
        type="list",
        formula1='"' + ",".join(GRADE_OPTIONS) + '"',
        allow_blank=True,
        error="建议选择“一年级”至“九年级”；其他写法可保留，未知可留空。",
        errorTitle="请核对年级",
        errorStyle="warning",
        prompt="年级选填，未知可留空；初一至初三会自动统一为七至九年级。",
        promptTitle="选择年级",
        showErrorMessage=True,
        showInputMessage=True,
    )
    worksheet.add_data_validation(branch_validation)
    branch_validation.add("I2:I2001")
    worksheet.add_data_validation(grade_validation)
    grade_validation.add("D2:D2001")
    for column in ("A", "D", "G", "I", "K"):
        worksheet.column_dimensions[column].number_format = "@"
    for index in range(1, len(ROSTER_TEMPLATE_HEADERS) + 1):
        letter = worksheet.cell(1, index).column_letter
        worksheet.column_dimensions[letter].width = 3 if not ROSTER_TEMPLATE_HEADERS[index - 1] else 18
    workbook.save(destination)
    return destination
