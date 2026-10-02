from pathlib import Path
import tempfile
import unittest
from unittest import mock
from zipfile import ZIP_DEFLATED, ZipFile

from app.models import (AnswerItem, ExamSession, ScoreSection, TemplatePackage,
                        ValidationError, validate_answer_key)
from app.answer_key_import import generate_answer_key_template, read_answer_key
from app.roster_import import (ROSTER_COLUMNS, ROSTER_TEMPLATE_HEADERS,
                               generate_roster_template, normalize_exam_id, read_roster,
                               save_roster_snapshot)
from app.review_store import (audit_entry, load_review_session,
                              save_review_session)
from app.scanner import (PART_SCORE_COLUMNS, RESULT_COLUMNS, RESULT_HEADER_LABELS, SUMMARY_SEPARATOR_COLUMN,
                         export_results, load_students, update_review_flag, update_score_totals)
from app.ui_helpers import (apply_exam_id, apply_mark, apply_student_supplement,
                            open_original_in_preview, prepare_mark_review)
from app.legacy_profile import answers_to_legacy_key, legacy_v1_package
from openpyxl import Workbook, load_workbook


def approved_template() -> TemplatePackage:
    return TemplatePackage(
        template_id="demo-card",
        name="演示答题卡",
        version="1.0",
        status="approved",
        score_sections=(ScoreSection("listening_p1", "Listening Part 1", 5),),
        answer_items=(AnswerItem("L01", "第 1 题", ("A", "B", "C"), "listening_p1"),),
    )


class ModelTests(unittest.TestCase):
    def test_s_student_number_normalization_matches_all_roster_headers(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            for header in ("学号", "学 号", "Exam ID", "考号", "准考证号"):
                with self.subTest(header=header):
                    source = Path(temp) / "名单.xlsx"
                    workbook = Workbook()
                    workbook.active.append([header, "中文名"])
                    workbook.active.append(["S10086", "测试学生"])
                    workbook.save(source)
                    imported = read_roster(source)
                    self.assertEqual(imported.rows[0]["Exam ID"], "010086")
                    self.assertEqual(imported.source_rows[0][0], "S10086")

    def test_s_and_numeric_aliases_cannot_create_two_students(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "重复名单.xlsx"
            workbook = Workbook()
            workbook.active.append(["Exam ID", "Chinese Name"])
            workbook.active.append(["S10086", "学生甲"])
            workbook.active.append(["010086", "学生乙"])
            workbook.save(source)
            with self.assertRaisesRegex(ValidationError, "重复"):
                read_roster(source)

    def test_manual_identity_confirmation_accepts_s_and_numeric_ids(self) -> None:
        student = {"Chinese Name": "测试学生", "Year Level": "六年级", "Branch": "测试分校",
                   "Class": "测试班级", "Exam Session": "测试场次", "Original Exam ID": "S10086"}
        for entered_id in ("010086", "S10086", " s10086 ", "10086"):
            with self.subTest(entered_id=entered_id):
                record = {"Source Image": "card.png", "Exam ID": "00?086", "Identity Issue": "考号识别异常",
                          **{column: 0 for column in PART_SCORE_COLUMNS}}
                item_rows = [{"Source Image": "card.png", "Exam ID": "00?086", "Section": "listening_part1", "Is Correct": "N"}]
                valid, _ = apply_exam_id(record, item_rows, entered_id, {"010086": student}, operator="测试老师")
                self.assertTrue(valid)
                self.assertEqual(record["Exam ID"], "010086")
                self.assertEqual(record["Original Exam ID"], "S10086")
                self.assertEqual(item_rows[0]["Exam ID"], "010086")
                self.assertEqual(record["Identity Issue"], "")

    def test_unknown_identity_is_rejected_without_modifying_record(self) -> None:
        record = {"Source Image": "card.png", "Exam ID": "00?086", "Identity Issue": "考号识别异常"}
        original = dict(record)
        valid, message = apply_exam_id(record, [], "S10086", {})
        self.assertFalse(valid)
        self.assertIn("010086", message)
        self.assertIn("后补学生", message)
        self.assertEqual(record, original)
        self.assertEqual(normalize_exam_id("S100860"), "S100860")

    def test_roster_snapshot_match_and_export_preserve_original_s_number(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for header in ("学号", "Exam ID", "考号"):
                with self.subTest(header=header):
                    source = root / "source.xlsx"
                    workbook = Workbook()
                    workbook.active.append([header, "姓名", "年级", "分校", "班级", "笔试时间"])
                    workbook.active.append(["S10086", "测试学生", "六年级", "测试分校", "测试班级", "测试场次"])
                    workbook.save(source)
                    config = root / header
                    save_roster_snapshot(source, config)
                    record = {"Source Image": "card.png", "File Name": "card.png", "Exam ID": "00?086",
                              "Identity Issue": "考号识别异常", **{column: 0 for column in PART_SCORE_COLUMNS}}
                    with mock.patch("app.scanner.CONFIG_DIR", config):
                        students = load_students()
                        self.assertEqual(students["010086"]["Original Exam ID"], "S10086")
                        valid, _ = apply_exam_id(record, [], "S10086", students)
                        self.assertTrue(valid)
                        result, _ = export_results([record], [], root, include_items=False)
                    exported = load_workbook(result, read_only=True, data_only=True)
                    rows = list(exported["成绩汇总"].iter_rows(values_only=True))
                    self.assertEqual(rows[1][list(rows[0]).index(header)], "S10086")
                    self.assertEqual(record["Exam ID"], "010086")
                    exported.close()

    def test_missing_student_s_identity_is_preserved_in_export(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "roster_source.xlsx"
            workbook = Workbook()
            workbook.active.append(["学号", "姓名", "年级", "分校", "班级", "笔试时间"])
            workbook.active.append(["S10086", "测试学生", "六年级", "测试分校", "测试班级", "测试场次"])
            workbook.save(source)
            record = {"Source Image": "missing.png", "Exam ID": "00?999", "Identity Issue": "考号识别异常",
                      **{column: 0 for column in PART_SCORE_COLUMNS}}
            apply_student_supplement(record, [], {"Chinese Name": "后补学生", "Year Level": "六年级",
                                     "Branch": "测试分校", "Class": "测试班级", "Exam Session": "测试场次"}, exam_id="S19999")
            with mock.patch("app.scanner.CONFIG_DIR", root):
                result, _ = export_results([record], [], root, include_items=False)
            exported = load_workbook(result, read_only=True, data_only=True)
            self.assertEqual(exported["成绩汇总"].cell(2, 1).value, "S19999")
            self.assertEqual(record["Exam ID"], "019999")
            exported.close()

    def test_unread_part2_and_part3_grader_scores_require_review_with_roi(self) -> None:
        template = {
            "score_part2": {"0": {"x": 1}, "1": {"x": 2}, "5": {"x": 3}},
            "score_part3": {"0": {"x": 4}, "9": {"x": 5}, "10": {"x": 6}},
        }
        blank = mock.Mock(value=None, ratios={"0": 0.02, "1": 0.03, "5": 0.04})
        ambiguous = mock.Mock(value=None, ratios={"0": 0.20, "9": 0.19, "10": 0.03})
        record = {"Score Entry States": {}}
        with (
            mock.patch("app.ui_helpers.QUESTION_GROUPS", []),
            mock.patch("app.ui_helpers.SCORE_FIELDS", {
                "score_part2": "Written_Part2", "score_part3": "Written_Part3",
            }),
            mock.patch("app.ui_helpers._read_image", return_value=object()),
            mock.patch("app.ui_helpers.normalize", return_value=object()),
            mock.patch("app.ui_helpers._read_one", side_effect=[blank, ambiguous]),
            mock.patch("app.ui_helpers._crop", side_effect=["part2-roi", "part3-roi"]),
        ):
            issues = prepare_mark_review(record, Path("card.jpg"), template)

        self.assertEqual([issue["key"] for issue in issues], ["score:score_part2", "score:score_part3"])
        self.assertTrue(all(issue["kind"] == "grader_score" for issue in issues))
        self.assertTrue(all(issue["allow_blank"] is False for issue in issues))
        self.assertEqual([issue["crop"] for issue in issues], ["part2-roi", "part3-roi"])
        self.assertEqual(issues[0]["choices"], ["0", "1", "5"])
        self.assertEqual(record["Score Entry States"]["Written_Part2"], "BLANK")
        self.assertEqual(record["Score Entry States"]["Written_Part3"], "AMBIGUOUS")

    def test_restoring_review_does_not_overwrite_saved_manual_grader_score(self) -> None:
        template = {"score_part2": {"0": {"x": 1}, "5": {"x": 2}}}
        blank = mock.Mock(value=None, ratios={"0": 0.02, "5": 0.03})
        record = {
            "Written_Part2": 5,
            "Score Entry States": {"Written_Part2": "MANUAL"},
            "Original Scores": {"Written_Part2": 0},
        }
        with (
            mock.patch("app.ui_helpers.QUESTION_GROUPS", []),
            mock.patch("app.ui_helpers.SCORE_FIELDS", {"score_part2": "Written_Part2"}),
            mock.patch("app.ui_helpers._read_image", return_value=object()),
            mock.patch("app.ui_helpers.normalize", return_value=object()),
            mock.patch("app.ui_helpers._read_one", return_value=blank),
            mock.patch("app.ui_helpers._crop", return_value="part2-roi"),
        ):
            issues = prepare_mark_review(record, Path("card.jpg"), template)

        self.assertEqual(len(issues), 0)
        self.assertEqual(record["Written_Part2"], 5)
        self.assertEqual(record["Score Entry States"]["Written_Part2"], "MANUAL")
        self.assertEqual(record["Original Score Entry States"]["Written_Part2"], "BLANK")
        self.assertIsNone(record["Original Scores"]["Written_Part2"])

    def test_grader_score_failure_is_not_reported_as_whole_part_blank(self) -> None:
        record = {
            "Source Image": "card.jpg",
            "Identity Issue": "",
            "Score Entry States": {
                "Written_Part2": "BLANK", "Written_Part3": "AMBIGUOUS", "Writing": "SCORED",
            },
        }
        update_review_flag(record, [])
        self.assertEqual(record["Needs Review"], "是")
        self.assertEqual(record["Review Reason"], "教师登分区未成功读取：Written Part 2、Written Part 3")
        self.assertEqual(record["Blank Parts"], "")
        self.assertNotIn("答题卡整Part未填写", record["Review Reason"])

    def test_machine_scored_zero_does_not_inherit_unrelated_reviewer(self) -> None:
        record = {
            "Source Image": "card.jpg", "File Name": "card.jpg", "Exam ID": "001001",
            "Chinese Name": "测试学生", "Year Level": "六年级", "Branch": "测试分校",
            "Class": "A1", "Exam Session": "9.24", "Identity Issue": "",
            "Written_Part2": 0, "Written_Part3": 0, "Writing": 0,
            "Score Entry States": {
                "Written_Part2": "SCORED", "Written_Part3": "SCORED", "Writing": "SCORED",
            },
            "Original Scores": {"Written_Part2": 0, "Written_Part3": 0, "Writing": 0},
            "Original Score Entry States": {
                "Written_Part2": "SCORED", "Written_Part3": "SCORED", "Writing": "SCORED",
            },
            "Last Operator": "其他题目的复核人",
            "Last Review Basis": "答题卡图片",
            "Last Changed At": "2026-09-24T02:00:00",
            **{column: 0 for column in (
                "Listening_Part1", "Listening_Part2", "Listening_Part3", "Written_Part1",
                "Written_Part4", "Written_Part5", "Written_Part6", "Written_Part7", "Written_Part8",
            )},
        }
        with tempfile.TemporaryDirectory() as temp:
            result, _ = export_results([record], [], Path(temp), "机器零分", include_items=True)
            workbook = load_workbook(result, read_only=True, data_only=True)
            detail = list(workbook["逐题明细"].iter_rows(values_only=True))
            header = {value: index for index, value in enumerate(detail[0])}
            part2 = next(row for row in detail[1:] if row[header["Part"]] == "Written Part 2")
            self.assertEqual(part2[header["最终生效答案"]], 0)
            self.assertEqual(part2[header["作答状态"]], "已填分")
            self.assertEqual(part2[header["是否人工修正"]], "否")
            self.assertIsNone(part2[header["核查依据"]])
            self.assertIsNone(part2[header["操作人"]])
            self.assertIsNone(part2[header["修改时间"]])
            workbook.close()

    def test_manual_grader_score_writes_final_score_and_audit_state(self) -> None:
        record = {
            "Source Image": "card.jpg",
            "File Name": "card.jpg",
            "Exam ID": "001001",
            "Chinese Name": "测试学生",
            "Year Level": "六年级",
            "Branch": "测试分校",
            "Class": "A1",
            "Exam Session": "9.24",
            "Identity Issue": "",
            "Written_Part2": 0,
            "Score Entry States": {"Written_Part2": "BLANK"},
            "Original Scores": {"Written_Part2": None, "Written_Part3": None},
            "Original Score Entry States": {"Written_Part2": "BLANK", "Written_Part3": "AMBIGUOUS"},
            **{column: 0 for column in (
                "Listening_Part1", "Listening_Part2", "Listening_Part3", "Written_Part1",
                "Written_Part3", "Written_Part4", "Written_Part5", "Written_Part6",
                "Written_Part7", "Written_Part8", "Writing",
            )},
        }
        issue = {"field": "score_part2", "kind": "grader_score", "title": "Written Part 2"}
        apply_mark(
            record, [], issue, "5", operator="张老师", review_basis="答题卡图片",
            changed_at="2026-09-24T01:30:00",
        )
        self.assertEqual(record["Written_Part2"], 5)
        self.assertEqual(record["Score Entry States"]["Written_Part2"], "MANUAL")
        self.assertEqual(record["Last Operator"], "张老师")
        self.assertEqual(record["Last Review Basis"], "答题卡图片")
        self.assertEqual(record["Total"], 5)
        self.assertEqual(record["Needs Review"], "否")

        part3_issue = {"field": "score_part3", "kind": "grader_score", "title": "Written Part 3"}
        apply_mark(
            record, [], part3_issue, "9", operator="张老师", review_basis="答题卡图片",
            changed_at="2026-09-24T01:31:00",
        )
        self.assertEqual(record["Written_Part3"], 9)
        self.assertEqual(record["Score Entry States"]["Written_Part3"], "MANUAL")
        self.assertEqual(record["Total"], 14)

        with tempfile.TemporaryDirectory() as temp:
            result, _ = export_results([record], [], Path(temp), "教师登分复核", include_items=True)
            workbook = load_workbook(result, read_only=True, data_only=True)
            detail = list(workbook["逐题明细"].iter_rows(values_only=True))
            header = {value: index for index, value in enumerate(detail[0])}
            rows_by_part = {row[header["Part"]]: row for row in detail[1:]}
            part2 = rows_by_part["Written Part 2"]
            part3 = rows_by_part["Written Part 3"]
            self.assertIsNone(part2[header["原始识别答案"]])
            self.assertEqual(part2[header["最终生效答案"]], 5)
            self.assertEqual(part2[header["原始作答状态"]], "教师登分区未成功读取")
            self.assertEqual(part2[header["作答状态"]], "已人工修正")
            self.assertEqual(part2[header["是否人工修正"]], "是")
            self.assertEqual(part2[header["操作人"]], "张老师")
            self.assertEqual(part3[header["最终生效答案"]], 9)
            self.assertEqual(part3[header["是否人工修正"]], "是")
            workbook.close()

    def test_isolated_blank_candidate_requires_single_question_review(self) -> None:
        template = {
            "written_objective": {
                "part7": {
                    "39": {"A": {}, "B": {}, "C": {}},
                    "40": {"A": {}, "B": {}, "C": {}},
                }
            }
        }
        blank = mock.Mock(value=None, ratios={"A": 0.05, "B": 0.07, "C": 0.04})
        answered = mock.Mock(value="A", ratios={"A": 0.40, "B": 0.05, "C": 0.04})
        with (
            mock.patch("app.ui_helpers.QUESTION_GROUPS", [("written_objective", "part7", "written_part7")]),
            mock.patch("app.ui_helpers.SCORE_FIELDS", {}),
            mock.patch("app.ui_helpers._read_image", return_value=object()),
            mock.patch("app.ui_helpers.normalize", return_value=object()),
            mock.patch("app.ui_helpers._read_one", side_effect=[blank, answered]),
            mock.patch("app.ui_helpers._crop", return_value=None),
        ):
            issues = prepare_mark_review({}, Path("card.jpg"), template)

        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0]["key"], "question:written_part7:39")
        self.assertTrue(issues[0]["blank_candidate"])
        self.assertEqual(issues[0]["detail"], "当前识别结果：空白候选")

    def test_whole_part_without_readable_answers_keeps_bulk_review(self) -> None:
        template = {
            "written_objective": {
                "part7": {
                    "39": {"A": {}, "B": {}, "C": {}},
                    "40": {"A": {}, "B": {}, "C": {}},
                }
            }
        }
        blank = mock.Mock(value=None, ratios={"A": 0.05, "B": 0.07, "C": 0.04})
        ambiguous = mock.Mock(value=None, ratios={"A": 0.24, "B": 0.21, "C": 0.04})
        with (
            mock.patch("app.ui_helpers.QUESTION_GROUPS", [("written_objective", "part7", "written_part7")]),
            mock.patch("app.ui_helpers.SCORE_FIELDS", {}),
            mock.patch("app.ui_helpers._read_image", return_value=object()),
            mock.patch("app.ui_helpers.normalize", return_value=object()),
            mock.patch("app.ui_helpers._read_one", side_effect=[blank, ambiguous]),
            mock.patch("app.ui_helpers._crop", return_value=None),
        ):
            issues = prepare_mark_review({}, Path("card.jpg"), template)

        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0]["kind"], "part_empty")
        self.assertEqual(len(issues[0]["question_issues"]), 2)

    def test_preview_opens_only_current_image_without_restoring_old_windows(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            image = Path(temp) / "图片2.png"
            image.touch()
            with mock.patch("app.ui_helpers.subprocess.run") as run:
                run.return_value.returncode = 0
                run.return_value.stderr = ""
                opened, _ = open_original_in_preview(image)
            self.assertTrue(opened)
            self.assertEqual(
                run.call_args.args[0],
                ["open", "-n", "-F", "-a", "Preview", str(image)],
            )

    def test_answer_key_must_match_template(self) -> None:
        template = approved_template()
        validate_answer_key(template.answer_items, {"L01": "B"})
        with self.assertRaises(ValidationError):
            validate_answer_key(template.answer_items, {"L01": "D"})

    def test_exam_requires_an_approved_template_and_local_inputs(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            roster = root / "students.xlsx"
            answer = root / "answers.xlsx"
            roster.touch()
            answer.touch()
            session = ExamSession("august-test", "八月定位测", approved_template(), roster, answer, root)
            session.validate_ready()

    def test_draft_template_cannot_be_used_for_exam(self) -> None:
        draft = TemplatePackage("draft-card", "草稿", "1.0", "draft", approved_template().answer_items, approved_template().score_sections)
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "students.xlsx").touch()
            (root / "answers.xlsx").touch()
            session = ExamSession("draft-test", "草稿考试", draft, root / "students.xlsx", root / "answers.xlsx", root)
            with self.assertRaises(ValidationError):
                session.validate_ready()

    def test_generated_answer_template_is_valid_after_teacher_entry(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            answer_file = Path(temp) / "答案录入.xlsx"
            generate_answer_key_template(approved_template(), answer_file)
            workbook = load_workbook(answer_file)
            workbook.active["E2"] = "C"
            workbook.save(answer_file)
            self.assertEqual(read_answer_key(answer_file, approved_template()), {"L01": "C"})

    def test_roster_rejects_duplicate_exam_id_and_saves_verified_snapshot(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            roster = root / "名单.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Students"
            worksheet.append(["Exam ID", "Name", "Class", "Branch"])
            worksheet.append(["261001", "学生甲", "A", "分校甲"])
            workbook.save(roster)
            self.assertEqual(read_roster(roster).count, 1)
            snapshot = save_roster_snapshot(roster, root / "exam")
            self.assertTrue(snapshot.is_file())
            worksheet.append(["261001", "学生乙", "B", "分校乙"])
            workbook.save(roster)
            with self.assertRaises(ValidationError):
                read_roster(roster)

    def test_roster_accepts_chinese_headers_on_non_first_sheet_and_normalizes_snapshot(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            roster = root / "参考成绩.xlsx"
            workbook = Workbook()
            workbook.active.title = "参考比例（动态表）"
            worksheet = workbook.create_sheet("目标学生名单")
            worksheet.append(["考号", "姓名", "班级", "上课校区", "Total"])
            worksheet.append(["261001", "学生甲", "A1", "分校甲", 72])
            workbook.save(roster)

            imported = read_roster(roster)
            self.assertEqual(imported.count, 1)
            self.assertEqual(imported.rows[0], {"Exam ID": "261001", "Chinese Name": "学生甲", "Year Level": "", "Branch": "分校甲", "Class": "A1", "Exam Session": ""})

            snapshot = save_roster_snapshot(roster, root / "exam")
            saved = load_workbook(snapshot, read_only=True, data_only=True)
            self.assertEqual(saved.active.title, "Students")
            self.assertEqual(tuple(next(saved.active.iter_rows(min_row=1, max_row=1, values_only=True))), ROSTER_COLUMNS)
            saved.close()

    def test_full_roster_template_and_chinese_headers(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            template = Path(temp) / "学生名单导入模板.xlsx"
            generate_roster_template(template)
            workbook = load_workbook(template)
            headers = tuple(cell.value or "" for cell in workbook.active[1])
            self.assertEqual(headers, ROSTER_TEMPLATE_HEADERS)
            self.assertEqual(workbook.active.title, "Sheet1")
            self.assertEqual(workbook.active.max_row, 1)
            validations = list(workbook.active.data_validations.dataValidation)
            self.assertEqual(len(validations), 2)
            self.assertTrue(any("牡丹广场分校" in str(item.formula1) for item in validations))
            self.assertTrue(any("七年级" in str(item.formula1) for item in validations))
            self.assertTrue(any("I2:I2001" in str(item.sqref) for item in validations))
            self.assertTrue(any("D2:D2001" in str(item.sqref) for item in validations))
            workbook.close()

    def test_original_roster_columns_are_preserved_before_score_columns(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "原始名单模板.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            original_headers = list(ROSTER_TEMPLATE_HEADERS)
            sheet.append(original_headers)
            sheet.append(["S00001", "张佳音Joy", "张佳音", "七年级", "英语", "RIB", "一班", "2026-12", "中心分校", "李老师", "第1批", "保留原备注"])
            workbook.save(source)
            config = root / "config"
            save_roster_snapshot(source, config, overwrite=True)
            imported = read_roster(config / "student_list.xlsx")
            self.assertEqual(imported.rows[0]["Chinese Name"], "张佳音")

            record = {
                "File Name": "000001.jpg", "Source Image": "000001.jpg",
                "Exam ID": "000001", "Chinese Name": "张佳音", "Year Level": "七年级",
                "Branch": "中心分校", "Class": "一班", "Exam Session": "第1批",
                "Identity Issue": "", "Score Entry States": {},
                **{column: 0 for column in (
                    "Listening_Part1", "Listening_Part2", "Listening_Part3", "Written_Part1",
                    "Written_Part2", "Written_Part3", "Written_Part4", "Written_Part5",
                    "Written_Part6", "Written_Part7", "Written_Part8", "Writing",
                )},
            }
            with mock.patch("app.scanner.CONFIG_DIR", config):
                result, _ = export_results([record], [], root / "out", "原名单导出")
            exported = load_workbook(result, read_only=True, data_only=True)
            summary = list(exported["成绩汇总"].iter_rows(values_only=True))
            self.assertEqual(
                [value or "" for value in summary[0][:len(original_headers)]],
                original_headers,
            )
            self.assertEqual(summary[1][0], "S00001")
            self.assertEqual(summary[1][1], "张佳音Joy")
            self.assertEqual(summary[1][11], "保留原备注")
            self.assertEqual(summary[0][26], None)
            self.assertEqual(summary[0][27], "File Name")
            self.assertEqual(summary[1][27], "000001.jpg")
            self.assertEqual(summary[0][28], "Listening_Part1")
            self.assertEqual(summary[1][28], 0)
            exported.close()

    def test_extra_cells_beyond_roster_headers_do_not_shift_handoff_columns(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "roster_source.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(ROSTER_TEMPLATE_HEADERS)
            sheet.append([
                "S10086", "张佳音Joy", "张佳音", "六年级", "RI", "IS12", "RIA",
                "2027", "新街分校", "李老师", "9.24", "", "否", "13800000000",
                "", "", "", "", "", "", "", "", "", "", "", "",
                # 表头范围之外残留的旧单元格，曾导致后续成绩列整体右移。
                "旧数据1", "旧数据2", "旧数据3",
            ])
            workbook.save(source)
            record = {
                "File Name": "010086.jpg", "Source Image": "010086.jpg", "Exam ID": "010086",
                "Chinese Name": "张佳音", "Year Level": "六年级", "Branch": "新街分校",
                "Class": "RIA", "Exam Session": "9.24", "Status": "OK",
                "Blank Parts": "", "Blank Questions": "", "Needs Review": "否",
                "Review Reason": "", "Identity Issue": "", "Total": 55,
                **{column: 0 for column in (
                    "Listening_Part1", "Listening_Part2", "Listening_Part3", "Written_Part1",
                    "Written_Part2", "Written_Part3", "Written_Part4", "Written_Part5",
                    "Written_Part6", "Written_Part7", "Written_Part8", "Writing",
                )},
            }
            with mock.patch("app.scanner.CONFIG_DIR", root):
                result, _ = export_results([record], [], root, "列对齐", include_items=False)
            exported = load_workbook(result, read_only=True, data_only=True)
            rows = list(exported["成绩汇总"].iter_rows(values_only=True))
            header = list(rows[0])
            values = rows[1]
            self.assertEqual(len(header), len(ROSTER_TEMPLATE_HEADERS) + 20)
            self.assertEqual(values[header.index("File Name")], "010086.jpg")
            self.assertEqual(values[header.index("需要人工核对学生试卷")], "否")
            self.assertIsNone(values[header.index("异常原因")])
            self.assertIsNone(values[header.index("身份异常")])
            self.assertNotIn("旧数据1", values)
            exported.close()

    def test_original_student_number_is_normalized_only_for_scanning(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "学业水平测试统计名单-模板.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(ROSTER_TEMPLATE_HEADERS)
            sheet.append([
                "S10086", "张佳音Joy", "张佳音", "六年级", "RI", "IS12", "RIA", "2026/9/4",
                "新街分校", "李美满Happy", "8.22（周六）", "", "", "130XXXXXXXX",
                "", "", "", "", "", "", "", "", "", "", "", "", "",
            ])
            workbook.save(source)
            imported = read_roster(source)
            self.assertEqual(imported.rows[0]["Exam ID"], "010086")
            self.assertEqual(imported.rows[0]["Chinese Name"], "张佳音")
            self.assertEqual(imported.source_rows[0][0], "S10086")

    def test_blank_chinese_name_cell_falls_back_to_bilingual_name(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "姓名公式未缓存.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(ROSTER_TEMPLATE_HEADERS)
            sheet.append([
                "S11339", "张景溪Echo", None, "六年级", "RI", "IS12", "RIA",
                "2027", "新街分校", "李老师", "8.22（周六）",
            ])
            workbook.save(source)

            imported = read_roster(source)

            self.assertEqual(imported.count, 1)
            self.assertEqual(imported.rows[0]["Exam ID"], "011339")
            self.assertEqual(imported.rows[0]["Chinese Name"], "张景溪")

    def test_written_test_time_header_is_used_as_exam_session(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "笔试时间名单.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            headers = list(ROSTER_TEMPLATE_HEADERS)
            headers[10] = "笔试时间"
            sheet.append(headers)
            sheet.append([
                "S11339", "张景溪Echo", "张景溪", "六年级", "RI", "IS12", "RIA",
                "2027", "新街分校", "李老师", "9.13 新增",
            ])
            workbook.save(source)

            imported = read_roster(source)

            self.assertEqual(imported.rows[0]["Exam Session"], "9.13 新增")

    def test_roster_ignores_incorrect_worksheet_dimension_metadata(self) -> None:
        """Excel 可以显示所有列，即使文件内部错把有效范围写成 A1:A1。"""
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source.xlsx"
            corrupted = root / "incorrect-dimension.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(ROSTER_TEMPLATE_HEADERS)
            sheet.append([
                "S10086", "张佳音Joy", "张佳音", "六年级", "RI", "IS12", "RIA", "2026/9/4",
                "新街分校", "李美满Happy", "8.22（周六）", "", "", "130XXXXXXXX",
                "", "", "", "", "", "", "", "", "", "", "", "", "",
            ])
            sheet.append([None, "陈尚钰Rock", "陈尚钰", "初一", None, None, "外部"])
            workbook.save(source)
            with ZipFile(source) as input_zip, ZipFile(corrupted, "w", ZIP_DEFLATED) as output_zip:
                for item in input_zip.infolist():
                    data = input_zip.read(item.filename)
                    if item.filename == "xl/worksheets/sheet1.xml":
                        original = data
                        data = data.replace(b'<dimension ref="A1:AA3" />', b'<dimension ref="A1:A1" />')
                        self.assertNotEqual(data, original)
                    output_zip.writestr(item, data)
            imported = read_roster(corrupted)
            self.assertEqual(imported.count, 1)
            self.assertEqual(imported.rows[0]["Exam ID"], "010086")
            self.assertEqual(imported.rows[0]["Chinese Name"], "张佳音")
            self.assertEqual(len(imported.source_headers), 26)
            self.assertEqual(imported.skipped_rows, ("第 3 行（陈尚钰）缺少学号",))

    def test_score_export_has_parts_total_blank_separator_and_blank_details(self) -> None:
        record = {
            "Listening_Part1": 5, "Listening_Part2": 5, "Listening_Part3": 5,
            "Written_Part1": 5, "Written_Part2": 5, "Written_Part3": 10,
            "Written_Part4": 5, "Written_Part5": 5, "Written_Part6": 5,
            "Written_Part7": 7, "Written_Part8": 8, "Writing": 15,
        }
        record.update({"File Name": "one.jpg", "Exam ID": "261001", "Chinese Name": "学生甲", "Year Level": "五年级", "Branch": "分校甲", "Class": "A1", "Exam Session": "第1批", "Status": "OK"})
        update_score_totals(record)
        self.assertEqual(record["Total"], 80)
        self.assertNotIn("Listening Total", RESULT_COLUMNS)
        self.assertNotIn("Written Total", RESULT_COLUMNS)
        total_index = RESULT_COLUMNS.index("Total")
        self.assertEqual(RESULT_COLUMNS[total_index + 1], SUMMARY_SEPARATOR_COLUMN)
        self.assertEqual(
            tuple(RESULT_COLUMNS[total_index + 2:]),
            ("Blank Parts", "Blank Questions", "Needs Review", "Review Reason", "Identity Issue"),
        )
        with tempfile.TemporaryDirectory() as temp:
            with mock.patch("app.scanner.CONFIG_DIR", Path(temp)):
                result, _ = export_results([record], [], Path(temp), "成绩", include_items=False)
            workbook = load_workbook(result, read_only=True, data_only=True)
            headers = tuple(next(workbook.active.iter_rows(min_row=1, max_row=1, values_only=True)))
            self.assertEqual(headers[1:7], ("考号", "中文名", "年级", "分校", "班级", "笔试时间"))
            self.assertNotIn("Listening Total", headers)
            self.assertNotIn("Written Total", headers)
            self.assertIsNone(headers[total_index + 1])
            self.assertEqual(headers[total_index + 2:total_index + 7], ("空白Part", "空白题号", "需要人工核对学生试卷", "异常原因", "身份异常"))
            workbook.close()

    def test_whole_blank_part_is_exported_for_manual_review(self) -> None:
        record = {column: 0 for column in (
            "Listening_Part1", "Listening_Part2", "Listening_Part3", "Written_Part1",
            "Written_Part2", "Written_Part3", "Written_Part4", "Written_Part5",
            "Written_Part6", "Written_Part7", "Written_Part8", "Writing",
        )}
        record.update({
            "File Name": "blank.jpg", "Source Image": "blank.jpg", "Exam ID": "261002",
            "Chinese Name": "学生乙", "Year Level": "五年级", "Branch": "分校乙",
            "Class": "A2", "Exam Session": "第2批", "Status": "OK",
            "Score Entry States": {"Written_Part2": "SCORED", "Written_Part3": "SCORED", "Writing": "SCORED"},
        })
        item_rows = [
            {"Exam ID": "261002", "Section": "listening_part1", "Question": str(number),
             "Marked Answer": "", "Correct Answer": "A", "Is Correct": "N",
             "Answer Status": "BLANK", "Source Image": "blank.jpg"}
            for number in range(1, 6)
        ]
        update_score_totals(record)
        update_review_flag(record, item_rows)
        self.assertEqual(record["Needs Review"], "是")
        self.assertEqual(record["Review Reason"], "答题卡整Part未填写")
        self.assertEqual(RESULT_HEADER_LABELS["Needs Review"], "需要人工核对学生试卷")
        self.assertEqual(RESULT_HEADER_LABELS["Review Reason"], "异常原因")
        self.assertEqual(record["Blank Parts"], "Listening Part 1")
        self.assertIn("1、2、3、4、5", record["Blank Questions"])

        with tempfile.TemporaryDirectory() as temp:
            result, item_path = export_results([record], item_rows, Path(temp), "空白复核", include_items=True)
            self.assertIsNone(item_path)
            workbook = load_workbook(result, read_only=True, data_only=True)
            self.assertEqual(workbook.sheetnames, ["成绩汇总", "逐题明细", "需复核名单", "信息待补充"])
            review_values = list(workbook["需复核名单"].iter_rows(values_only=True))
            review_header = {value: index for index, value in enumerate(review_values[0])}
            self.assertEqual(review_values[1][review_header["考号"]], "261002")
            self.assertEqual(review_values[1][review_header["空白Part"]], "Listening Part 1")
            self.assertEqual(review_values[1][review_header["需要人工核对学生试卷"]], "是")
            self.assertEqual(review_values[1][review_header["异常原因"]], "整Part未填涂，需复查原试卷")
            self.assertEqual(review_values[1][review_header["处理状态"]], "待复查原卷")
            self.assertEqual(review_values[1][review_header["说明"]], "本次学生整Part未填涂，建议复查原试卷。")
            self.assertNotIn("人工复核结论", review_header)
            self.assertNotIn("操作人", review_header)
            detail_values = list(workbook["逐题明细"].iter_rows(values_only=True))
            detail_header = {value: index for index, value in enumerate(detail_values[0])}
            self.assertEqual(detail_values[1][detail_header["作答状态"]], "空白")
            self.assertEqual(detail_values[1][detail_header["Part状态"]], "整Part空白")
            workbook.close()

            # 图片中已确认整 Part 空白后，系统内的修正已完成，
            # 但“需查原试卷”仍是 Sheet2 无法解决的后续事项。
            record["Reviewed Blank Sections"] = ["listening_part1"]
            update_review_flag(record, item_rows)
            self.assertEqual(record["Needs Review"], "否")
            audit = [audit_entry(
                record, operator="张老师", change_type="人工确认", target="Listening Part 1",
                original_value="待确认", corrected_value="整Part为空", review_basis="答题卡图片",
            )]
            resolved_result, _ = export_results(
                [record], item_rows, Path(temp), "空白复核已确认", include_items=True, audit_log=audit,
            )
            resolved_workbook = load_workbook(resolved_result, read_only=True, data_only=True)
            resolved_review = list(resolved_workbook["需复核名单"].iter_rows(values_only=True))
            self.assertEqual(len(resolved_review), 2)
            self.assertEqual(resolved_review[1][resolved_review[0].index("处理状态")], "待复查原卷")
            resolved_workbook.close()

    def test_manual_correction_preserves_original_and_exports_audit(self) -> None:
        record = {
            "File Name": "one.jpg", "Source Image": "one.jpg", "Exam ID": "261001",
            "Chinese Name": "学生甲", "Year Level": "五年级", "Branch": "分校甲",
            "Class": "A1", "Exam Session": "第1批", "Identity Issue": "",
            **{column: 0 for column in RESULT_COLUMNS if column.endswith(("Part1", "Part2", "Part3", "Part4", "Part5", "Part6", "Part7", "Part8"))},
            "Writing": 0,
        }
        item_rows = [{
            "Exam ID": "261001", "Section": "listening_part1", "Question": "1",
            "Original Marked Answer": "A", "Marked Answer": "A", "Correct Answer": "B",
            "Original Item Score": 0, "Is Correct": "N", "Original Answer Status": "INCORRECT",
            "Answer Status": "INCORRECT", "Manual Correction": "否", "Source Image": "one.jpg",
        }]
        issue = {"section": "listening_part1", "number": "1", "title": "题号 1"}
        with mock.patch("app.ui_helpers.load_answer_key", return_value={"listening_part1": {"1": "B"}}):
            apply_mark(record, item_rows, issue, "B", operator="张老师", review_basis="答题卡图片", changed_at="2026-09-19T10:00:00")
        self.assertEqual(item_rows[0]["Original Marked Answer"], "A")
        self.assertEqual(item_rows[0]["Marked Answer"], "B")
        self.assertEqual(item_rows[0]["Manual Correction"], "是")
        self.assertEqual(item_rows[0]["Operator"], "张老师")

        audit = [audit_entry(
            record, operator="张老师", change_type="修正填涂", target="题号 1",
            original_value="A", corrected_value="B", review_basis="答题卡图片",
        )]
        with tempfile.TemporaryDirectory() as temp:
            result, _ = export_results([record], item_rows, Path(temp), "修正测试", audit_log=audit)
            workbook = load_workbook(result, read_only=True, data_only=True)
            detail = list(workbook["逐题明细"].iter_rows(values_only=True))
            header = {value: index for index, value in enumerate(detail[0])}
            self.assertEqual(detail[1][header["原始识别答案"]], "A")
            self.assertEqual(detail[1][header["最终生效答案"]], "B")
            self.assertEqual(detail[1][header["是否人工修正"]], "是")
            self.assertEqual(detail[1][header["操作人"]], "张老师")
            review = list(workbook["需复核名单"].iter_rows(values_only=True))
            review_header = {value: index for index, value in enumerate(review[0])}
            self.assertEqual(len(review), 1)
            self.assertNotIn("人工复核结论", review_header)
            self.assertNotIn("核查依据", review_header)
            workbook.close()

    def test_unmatched_exam_id_is_kept_and_sorted_last(self) -> None:
        base_scores = {column: 0 for column in (
            "Listening_Part1", "Listening_Part2", "Listening_Part3", "Written_Part1",
            "Written_Part2", "Written_Part3", "Written_Part4", "Written_Part5",
            "Written_Part6", "Written_Part7", "Written_Part8", "Writing",
        )}
        matched = {**base_scores, "File Name": "a.jpg", "Source Image": "a.jpg", "Exam ID": "261001", "Chinese Name": "学生甲", "Identity Issue": ""}
        unmatched = {**base_scores, "File Name": "b.jpg", "Source Image": "b.jpg", "Exam ID": "999999", "Chinese Name": "", "Identity Issue": "考号不在名单中"}
        with tempfile.TemporaryDirectory() as temp:
            with mock.patch("app.scanner.CONFIG_DIR", Path(temp)):
                result, _ = export_results([unmatched, matched], [], Path(temp), "名单外考号")
            workbook = load_workbook(result, read_only=True, data_only=True)
            summary = list(workbook["成绩汇总"].iter_rows(values_only=True))
            header = {value: index for index, value in enumerate(summary[0])}
            self.assertEqual(summary[1][header["考号"]], "261001")
            self.assertEqual(summary[2][header["考号"]], "999999")
            self.assertIsNone(summary[2][header["中文名"]])
            self.assertEqual(summary[2][header["身份异常"]], "考号不在名单中")
            workbook.close()

    def test_unmatched_exam_id_can_be_confirmed_as_roster_omission(self) -> None:
        record = {
            "File Name": "missing.jpg", "Source Image": "missing.jpg",
            "Exam ID": "999999", "Chinese Name": "",
            "Identity Issue": "考号不在名单中",
            **{column: 0 for column in (
                "Listening_Part1", "Listening_Part2", "Listening_Part3",
                "Written_Part1", "Written_Part2", "Written_Part3", "Written_Part4",
                "Written_Part5", "Written_Part6", "Written_Part7", "Written_Part8", "Writing",
            )},
        }
        apply_student_supplement(
            record, [], {
                "Chinese Name": "学生乙", "Year Level": "五年级", "Branch": "分校乙",
                "Class": "B1", "Exam Session": "第2批",
            },
            operator="李老师", review_basis="纸质原卷", changed_at="2026-09-20T10:00:00",
        )
        self.assertEqual(record["Chinese Name"], "学生乙")
        self.assertEqual(record["Identity Issue"], "")
        self.assertTrue(record["Identity Confirmed"])
        self.assertEqual(record["Last Operator"], "李老师")
        self.assertEqual(record["Needs Review"], "否")

    def test_unread_exam_id_can_be_entered_while_supplementing_roster(self) -> None:
        record = {
            "Source Image": "unknown.jpg", "Exam ID": "00?999", "Identity Issue": "考号识别异常",
            **{column: 0 for column in (
                "Listening_Part1", "Listening_Part2", "Listening_Part3",
                "Written_Part1", "Written_Part2", "Written_Part3", "Written_Part4",
                "Written_Part5", "Written_Part6", "Written_Part7", "Written_Part8", "Writing",
            )},
        }
        item_rows = [{"Source Image": "unknown.jpg", "Exam ID": "00?999", "Section": "listening_part1", "Is Correct": "N"}]
        apply_student_supplement(
            record, item_rows, {
                "Chinese Name": "学生丙", "Year Level": "五年级", "Branch": "分校丙",
                "Class": "C1", "Exam Session": "第3批",
            },
            exam_id="005953", operator="王老师",
        )
        self.assertEqual(record["Exam ID"], "005953")
        self.assertEqual(item_rows[0]["Exam ID"], "005953")
        self.assertEqual(record["Identity Issue"], "")

    def test_review_session_round_trip(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            source_folder = Path(temp) / "images"
            source_folder.mkdir()
            with mock.patch("app.review_store.SESSION_DIR", Path(temp) / "sessions"):
                save_review_session(
                    source_folder,
                    [{"Source Image": "one.jpg", "Exam ID": "261001"}],
                    [{"Source Image": "one.jpg", "Question": "1"}],
                    [{"Source Image": "one.jpg", "Operator": "张老师"}],
                    {"one.jpg::question:1"}, {"one.jpg"}, "张老师",
                )
                restored = load_review_session(source_folder)
            self.assertIsNotNone(restored)
            self.assertEqual(restored["reviewer"], "张老师")
            self.assertEqual(restored["resolved"], {"one.jpg::question:1"})

    def test_explicit_zero_score_does_not_trigger_review(self) -> None:
        record = {
            "Source Image": "zero.jpg", "Score Entry States": {
                "Written_Part2": "SCORED", "Written_Part3": "SCORED", "Writing": "SCORED"
            }
        }
        item_rows = [{
            "Section": "listening_part1", "Question": "1", "Marked Answer": "B",
            "Correct Answer": "A", "Is Correct": "N", "Answer Status": "INCORRECT",
            "Source Image": "zero.jpg",
        }]
        update_review_flag(record, item_rows)
        self.assertEqual(record["Needs Review"], "否")

    def test_legacy_profile_round_trips_answers(self) -> None:
        template = {
            "listening": {"part1": {"1": {"A": {}, "B": {}}}, "part2": {}, "part3": {}},
            "written_objective": {"part1": {}, "part4": {}, "part5": {}, "part6": {}, "part7": {}, "part8": {}},
        }
        package = legacy_v1_package(template)
        self.assertEqual(package.answer_items[0].item_id, "listening_part1:1")
        self.assertEqual(answers_to_legacy_key({"listening_part1:1": "B"}), {"listening_part1": {"1": "B"}})


if __name__ == "__main__":
    unittest.main()
