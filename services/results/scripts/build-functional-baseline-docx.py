from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "考试结果查询系统_功能布局与客户体验基线_2026-09-07.md"
OUTPUT = ROOT / "docs" / "考试结果查询系统_功能布局与客户体验基线_2026-09-07.docx"


def set_run_font(run, name="Arial Unicode MS", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="5"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margin(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=(100, 100, 100))
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=8.5, color=(100, 100, 100))


def add_inline_runs(paragraph, text, size=10.5, color=(40, 40, 40)):
    parts = re.split(r"(《[^》]+》|\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        if bold:
            part = part[2:-2]
        elif code:
            part = part[1:-1]
        run = paragraph.add_run(part)
        set_run_font(run, name="Menlo" if code else "Arial Unicode MS", size=size, bold=bold, color=color)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(2.05)
    section.right_margin = Cm(2.05)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(40, 40, 40)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    style_specs = {
        "Title": (24, 0, 18),
        "Heading 1": (16, 15, 8),
        "Heading 2": (12.5, 10, 5),
    }
    for style_name, (size, before, after) in style_specs.items():
        style = doc.styles[style_name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "考试结果查询系统功能布局与客户体验基线    "
    for run in p.runs:
        set_run_font(run, size=8.5, color=(100, 100, 100))
    add_page_field(p)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append([cell.strip() for cell in lines[i].strip().strip("|").split("|")])
        i += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows, i


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_idx, values in enumerate(rows):
        row = table.rows[r_idx]
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx in range(cols):
            cell = row.cells[c_idx]
            value = values[c_idx] if c_idx < len(values) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margin(cell)
            if r_idx == 0:
                set_cell_shading(cell, "244C3F")
            elif r_idx % 2 == 0:
                set_cell_shading(cell, "F1F6F3")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 or c_idx == cols - 1 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_run_font(run, size=9.5, bold=r_idx == 0, color=(255, 255, 255) if r_idx == 0 else (35, 35, 35))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    i = 0
    first_title = True
    while i < len(lines):
        raw = lines[i].rstrip()
        text = raw.strip()
        if not text:
            i += 1
            continue
        if text.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if text.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline_runs(p, text[2:], size=24, color=(0, 0, 0))
            first_title = False
        elif text.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline_runs(p, text[3:], size=16, color=(0, 0, 0))
        elif text.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline_runs(p, text[4:], size=12.5, color=(0, 0, 0))
        elif re.match(r"^- ", text):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_inline_runs(p, text[2:])
        elif re.match(r"^\d+\. ", text):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            add_inline_runs(p, re.sub(r"^\d+\. ", "", text))
        elif text.startswith("版本日期：") or text.startswith("用途："):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            add_inline_runs(p, text.rstrip("  "), size=9.5, color=(95, 95, 95))
        else:
            p = doc.add_paragraph()
            add_inline_runs(p, text.rstrip("  "))
        i += 1

    core = doc.core_properties
    core.title = "考试结果查询系统功能布局与客户体验基线"
    core.subject = "系统功能 页面布局 家长体验 下一轮需求讨论"
    core.author = "项目工作组"
    core.keywords = "考试结果查询 微信小程序 成绩分析 六维评价 产品需求"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
